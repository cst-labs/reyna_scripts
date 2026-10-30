"""
ct_site_conflict_review.py
Reviews how well brainbodysep.tsv's static Order Type -> Site lookup matches
the ground-truth "Performing Location" field embedded in each imaging row's
Ordering Question Responses free text (e.g. "Performing Location: Tan Tock
Seng Hospital" / "... National Neuroscience Institute").

Three angles on the same underlying question -- does a static per-type
lookup actually hold up against what really happened?

  1. Conflict_Groups: every CSN where clean_ct_orders.py's own overlapping-
     exam-window grouping produces a cross-row NNI/DR site conflict (the
     group is NOT collapsed into one merged row because it spans both
     sites) -- one row per original order, with its tsv Site and its real
     Performing Location, side by side.
  2. Conflict_Types_Dedup: the distinct CT order types that show up inside
     those conflict groups, deduplicated, with each type's tsv Site and its
     Performing Location tally across the WHOLE dataset (not just the
     conflict rows) -- the larger sample is what actually tells you whether
     a type is reliably single-site or not.
  3. Mixed_Types_Temporal: every order type in brainbodysep.tsv whose real
     Performing Location is internally mixed (some rows NNI, some DR) --
     with the NNI/DR counts and an automatic classification of whether the
     mix is a clean temporal regime change, a scattered/intermittent mix,
     or a single stray outlier not worth treating as a pattern.

Outputs: analysis/ct_site_conflict_review.xlsx
"""
import os
import math
import re
from datetime import datetime, timedelta, date, time as time_type

import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
MERGED = os.path.join(BASE, "merged")
OUT = os.path.join(BASE, "analysis")
os.makedirs(OUT, exist_ok=True)


# ---------------------------------------------------------------------------
# Shared helpers (same implementations as clean_ct_orders.py / other scripts)
# ---------------------------------------------------------------------------
def is_blank(v):
    if v is None:
        return True
    try:
        if isinstance(v, float) and math.isnan(v):
            return True
    except Exception:
        pass
    if isinstance(v, str) and v.strip() == "":
        return True
    return False


def load_sheet(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
    seen = {}
    uh = []
    for h in headers:
        if h in seen:
            seen[h] += 1
            uh.append(f"{h}_{seen[h]}")
        else:
            seen[h] = 0
            uh.append(h)
    return pd.DataFrame(rows[1:], columns=uh)


def to_dt(val):
    if val is None:
        return None
    try:
        if isinstance(val, float) and math.isnan(val):
            return None
    except Exception:
        pass
    if isinstance(val, datetime):
        return val
    if isinstance(val, date):
        return datetime(val.year, val.month, val.day)
    if isinstance(val, str):
        s = val.strip()
        if not s:
            return None
        for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d",
                    "%d/%m/%Y %H:%M:%S", "%d-%m-%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(s, fmt)
            except ValueError:
                continue
        return None
    try:
        f = float(val)
        if math.isnan(f) or math.isinf(f):
            return None
        return datetime(1899, 12, 30) + timedelta(days=f)
    except Exception:
        return None


def to_time(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.time()
    if isinstance(val, time_type):
        return val
    if isinstance(val, str):
        s = val.strip()
        if not s:
            return None
        for fmt in ("%H:%M:%S", "%H:%M", "%I:%M:%S %p", "%I:%M %p"):
            try:
                return datetime.strptime(s, fmt).time()
            except ValueError:
                continue
        return None
    try:
        f = float(val)
        if math.isnan(f) or math.isinf(f):
            return None
        if 0 <= f < 2:
            ts = round(f * 86400)
            h, r = divmod(ts, 3600)
            m, s = divmod(r, 60)
            return time_type(h % 24, m, s)
    except Exception:
        pass
    return None


def combine_dt(d_val, t_val):
    d = to_dt(d_val)
    if d is None:
        return None
    if t_val is None:
        return d
    t = to_time(t_val)
    if t is None:
        return d
    return datetime(d.year, d.month, d.day, t.hour, t.minute, t.second)


def auto_col_width(ws, extra=2):
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + extra, 50)


LOC_TO_SITE = {
    "National Neuroscience Institute": "NNI",
    "Tan Tock Seng Hospital": "DR",
}


def extract_location(oqr):
    if is_blank(oqr):
        return None
    m = re.search(r"Performing Location:\s*([^\n]+)", str(oqr))
    return m.group(1).strip() if m else None


# ---------------------------------------------------------------------------
# Load imaging data + real Performing Location + tsv Site mapping
# ---------------------------------------------------------------------------
print("Loading merged/ct_orders_combined.xlsx ...")
df = load_sheet(os.path.join(MERGED, "ct_orders_combined.xlsx"), "Imaging")
print(f"  Loaded {len(df):,} imaging rows")

df["_perf_loc"] = df["Ordering Question Responses"].apply(extract_location)
df["_perf_site"] = df["_perf_loc"].map(LOC_TO_SITE)
df["exam_start_dt"] = df.apply(lambda r: combine_dt(r["Exam Start Date"], r["Exam Start Time"]), axis=1)
df["exam_end_dt"] = df.apply(lambda r: combine_dt(r["Exam End Date"], r["Exam End Time"]), axis=1)

print("Loading brainbodysep.tsv for CT site mapping...")
site_map_df = pd.read_csv(os.path.join(BASE, "brainbodysep.tsv"), sep="\t")
SITE_MAP = dict(zip(site_map_df["Order Type"].str.strip(), site_map_df["Site of Scan"].str.strip()))


def row_sites(desc):
    """A row's own tsv site(s) -- Order Description can be '; '-joined if
    the underlying order text itself lists multiple types."""
    if is_blank(desc):
        return set()
    parts = str(desc).split("; ") if "; " in str(desc) else [str(desc)]
    sites = set()
    for p in parts:
        s = SITE_MAP.get(p.strip())
        if s:
            sites.add(s)
    return sites


df["_sites"] = df["Order Description"].apply(row_sites)

valid = df[df["exam_start_dt"].notna() & df["exam_end_dt"].notna()].copy()


# ---------------------------------------------------------------------------
# Sheet 1: Conflict_Groups -- same transitive-overlap grouping per CSN as
# clean_ct_orders.py, keeping only groups whose member rows' tsv Site union
# spans BOTH NNI and DR (i.e. groups clean_ct_orders.py does NOT fully merge).
# ---------------------------------------------------------------------------
print("Grouping overlapping/identical-window imaging orders per CSN...")
groups_info = []
for csn, g in valid.groupby("CSN"):
    g = g.sort_values("exam_start_dt", kind="stable").reset_index()
    cur_group = []
    cur_max_end = None
    for _, r in g.iterrows():
        if cur_group and r["exam_start_dt"] <= cur_max_end:
            cur_group.append(r)
            cur_max_end = max(cur_max_end, r["exam_end_dt"])
        else:
            if len(cur_group) > 1:
                groups_info.append((csn, cur_group))
            cur_group = [r]
            cur_max_end = r["exam_end_dt"]
    if len(cur_group) > 1:
        groups_info.append((csn, cur_group))

conflict_rows = []
gid = 0
for csn, members in groups_info:
    union_sites = set()
    for m in members:
        union_sites |= m["_sites"]
    if {"NNI", "DR"}.issubset(union_sites):
        gid += 1
        for m in members:
            conflict_rows.append({
                "Group": gid,
                "CSN": csn,
                "Order ID": m.get("Order ID"),
                "Exam Start": m["exam_start_dt"],
                "Exam End": m["exam_end_dt"],
                "Order Description": m["Order Description"],
                "TSV Site(s)": ",".join(sorted(m["_sites"])) if m["_sites"] else "-",
                "Performing Location (raw)": m["_perf_loc"],
                "Performing Location Site": m["_perf_site"],
            })

conflict_groups = pd.DataFrame(conflict_rows).sort_values(["Group", "Exam Start"]).reset_index(drop=True)
print(f"  Conflicting groups (span both NNI and DR): {conflict_groups['Group'].nunique()} groups "
      f"/ {len(conflict_groups)} rows")


# ---------------------------------------------------------------------------
# Whole-dataset Performing Location tally per order type (used by both
# Sheet 2 and Sheet 3)
# ---------------------------------------------------------------------------
type_rows = []
for order_type, tsv_site in SITE_MAP.items():
    sub = df[df["Order Description"] == order_type]
    n = len(sub)
    n_with_loc = int(sub["_perf_site"].notna().sum())
    n_nni = int((sub["_perf_site"] == "NNI").sum())
    n_dr = int((sub["_perf_site"] == "DR").sum())
    if n_with_loc == 0:
        majority = None
    elif n_nni > n_dr:
        majority = "NNI"
    elif n_dr > n_nni:
        majority = "DR"
    else:
        majority = "TIE"
    type_rows.append({
        "Order Type": order_type, "TSV Site": tsv_site, "N (all rows)": n,
        "N w/ Performing Location": n_with_loc, "N NNI": n_nni, "N DR": n_dr,
        "Majority Performing Location Site": majority,
        "Agrees w/ TSV": (majority == tsv_site) if majority not in (None, "TIE") else None,
        "Internally Mixed": n_nni > 0 and n_dr > 0,
    })
type_tally = pd.DataFrame(type_rows)


# ---------------------------------------------------------------------------
# Sheet 2: Conflict_Types_Dedup -- distinct real CT types (tsv-mapped, i.e.
# excludes XR/other non-CT rows that only entered a conflict group because
# they overlapped a CT exam window) seen inside the 32 conflict groups,
# with each type's whole-dataset Performing Location tally.
# ---------------------------------------------------------------------------
conflict_ct_types = sorted(
    conflict_groups.loc[conflict_groups["TSV Site(s)"] != "-", "Order Description"].unique()
)
conflict_types_dedup = type_tally[type_tally["Order Type"].isin(conflict_ct_types)].copy()
conflict_types_dedup = conflict_types_dedup.sort_values("Order Type").reset_index(drop=True)
print(f"  Distinct CT types inside conflict groups: {len(conflict_types_dedup)}")


# ---------------------------------------------------------------------------
# Sheet 3: Mixed_Types_Temporal -- every order type with an internally
# mixed Performing Location (whole dataset), classified by whether the mix
# is a clean temporal regime change, a scattered/intermittent mix, or a
# single stray outlier.
# ---------------------------------------------------------------------------
def classify_temporal(order_type):
    sub = df[(df["Order Description"] == order_type) & df["_perf_site"].notna()].copy()
    sub["_month"] = sub.apply(lambda r: to_dt(r.get("Parent Order Date")), axis=1)
    sub = sub[sub["_month"].notna()]
    sub["_month"] = sub["_month"].apply(lambda d: (d.year, d.month))
    n_nni, n_dr = int((sub["_perf_site"] == "NNI").sum()), int((sub["_perf_site"] == "DR").sum())
    minority = min(n_nni, n_dr)

    monthly = sub.groupby(["_month", "_perf_site"]).size().unstack(fill_value=0)
    monthly = monthly.reindex(sorted(monthly.index))
    monthly_str = "; ".join(
        f"{y}-{m:02d}: NNI={row.get('NNI', 0)}/DR={row.get('DR', 0)}"
        for (y, m), row in monthly.iterrows()
    )

    if minority <= 1:
        return "Isolated outlier (single stray row, not a real pattern)", monthly_str

    majority_by_month = monthly.idxmax(axis=1)
    seq = list(majority_by_month)
    collapsed = [seq[0]]
    for s in seq[1:]:
        if s != collapsed[-1]:
            collapsed.append(s)
    if len(collapsed) == 2 and len(set(collapsed)) == 2:
        change_idx = next(i for i in range(1, len(seq)) if seq[i] != seq[i - 1])
        change_month = list(majority_by_month.index)[change_idx]
        return (f"Regime change: {collapsed[0]}-majority through "
                f"{list(majority_by_month.index)[change_idx - 1][0]}-{list(majority_by_month.index)[change_idx - 1][1]:02d}, "
                f"then {collapsed[1]} every month from {change_month[0]}-{change_month[1]:02d} onward "
                f"(no reversion)"), monthly_str
    return "Intermittent / mixed without a clean before-after split", monthly_str


mixed_types = type_tally[type_tally["Internally Mixed"]].copy()
classifications = mixed_types["Order Type"].apply(classify_temporal)
mixed_types["Temporal Classification"] = classifications.apply(lambda t: t[0])
mixed_types["Monthly NNI/DR Breakdown"] = classifications.apply(lambda t: t[1])
mixed_types = mixed_types.sort_values("Order Type").reset_index(drop=True)
print(f"  Order types with internally mixed Performing Location: {len(mixed_types)}")
for _, r in mixed_types.iterrows():
    print(f"    {r['Order Type']}: NNI={r['N NNI']} DR={r['N DR']} -> {r['Temporal Classification']}")


# ---------------------------------------------------------------------------
# Write
# ---------------------------------------------------------------------------
out_path = os.path.join(OUT, "ct_site_conflict_review.xlsx")
print(f"Writing {out_path} ...")
with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
    conflict_groups.to_excel(writer, sheet_name="Conflict_Groups", index=False)
    conflict_types_dedup.to_excel(writer, sheet_name="Conflict_Types_Dedup", index=False)
    mixed_types.to_excel(writer, sheet_name="Mixed_Types_Temporal", index=False)
    for sheet in writer.sheets.values():
        auto_col_width(sheet)

print("Done.")
