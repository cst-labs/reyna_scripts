"""
nni_llm_queue_slot_consistency_gemma_wrong_entity_cases.py
Pulls the full reasoning transcripts for the specific gemma4:e4b tool-calling
failure mode found by manual review of the run_index==1 ("first run") cases:
the model runs a pairwise compare_priority/compare_wait_time call, the *other*
item wins (i.e. that item should be scanned before new_order), and instead of
continuing to determine new_order's own rank, it reports the winning item's
position (almost always 1) as the final answer for new_order -- in several
cases explicitly writing "new_order is not Position 1" and then submitting
Position 1 anyway because it feels bound by the "reply with ONLY 'Position: N'"
instruction.

11 cases match this pattern out of 250 gemma4:e4b/tools run_index==1 cases
(16 non-PASS total there): 8 "clean" instances where the final answer is
literally the winning item's implied rank, and 3 "muddled" instances that
share the identical root cause (explicit "not Position 1" admission followed
by a guessed small number) without a clean single winner-substitution.

Reads:
  analysis/nni_llm_queue_slot_consistency_results.xlsx (Runs, Tool_Calls)
  analysis/nni_llm_queue_slot_consistency_raw.jsonl     (checkpoint -- source of
                                                          the literal prompt text,
                                                          which isn't carried into
                                                          the Runs sheet above)
Outputs: analysis/nni_llm_queue_slot_consistency_gemma_wrong_entity_cases.xlsx
"""
import json
import os

import pandas as pd
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ANALYSIS = os.path.join(BASE, "analysis")
RESULTS_PATH = os.path.join(ANALYSIS, "nni_llm_queue_slot_consistency_results.xlsx")
JSONL_PATH = os.path.join(ANALYSIS, "nni_llm_queue_slot_consistency_raw.jsonl")
OUT_PATH = os.path.join(ANALYSIS, "nni_llm_queue_slot_consistency_gemma_wrong_entity_cases.xlsx")

# (order_id, run_index, match_type, why_flagged) for the 11 cases identified by
# manual review of every gemma4:e4b/tools run_index==1 non-PASS transcript.
FLAGGED = [
    ("1783635096", 1, "clean",
     "compare_priority(new_order, item_1) -> item_1 wins (P2 beats Unknown); "
     "model still answers Position 1 for new_order."),
    ("1795981484", 1, "clean",
     "item_2 beats item_1 and beats new_order on priority; thinking concludes "
     "'item_2 must be the item scanned next... Conclusion: Position 1'."),
    ("1752206437", 1, "clean",
     "item_1 beats item_2 and new_order on wait time; thinking concludes "
     "'item_1 is the next item to be scanned. This is Position 1'."),
    ("1752723000", 1, "clean",
     "item_2 beats item_3 and new_order on wait time; thinking concludes "
     "'item_2 should be scanned next (Position 1)'."),
    ("1759915468; 1759915476", 1, "clean",
     "item_4 ties on priority then wins the wait-time tiebreak over new_order; "
     "thinking concludes 'item_4 is the next item to be scanned (Position 1)'."),
    ("1799314676", 1, "muddled",
     "Full tournament finds item_4 is the overall P1 winner (7 tool calls); "
     "thinking explicitly states 'new_order is not Position 1' then guesses "
     "Position 2 anyway ('I will follow the pattern of the prompt structure')."),
    ("1766112574", 1, "muddled",
     "item_3 ties then wins the wait-time tiebreak over new_order; thinking "
     "states 'new_order is not Position 1' then guesses Position 6 as 'the last "
     "item evaluated'."),
    ("1776020229; 1776020622; 1776020679; 1776020717", 1, "clean",
     "item_1 beats item_7, item_8, and new_order across 6 tool calls; thinking "
     "explicitly says 'I will output the position of the item I determined to "
     "be first' -- Position 1."),
    ("1779462640", 1, "clean",
     "item_9 ties then wins the wait-time tiebreak over new_order; thinking "
     "concludes 'item_9 is in Position 1'."),
    ("1747881172", 1, "clean",
     "item_3 beats new_order on priority; thinking states 'new_order is not "
     "Position 1... I will answer based on determination of the actual next "
     "item. The next item is item_3. Position: 1'."),
    ("1773150977", 1, "muddled",
     "item_7 and item_8 both beat new_order on priority; thinking states "
     "'Conclusion: new_order is not Position 1' then outputs Position 2 anyway "
     "('I must stop and trust the established hierarchy')."),
]


def auto_col_width(ws, extra=2, cap=80):
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + extra, cap)


runs = pd.read_excel(RESULTS_PATH, sheet_name="Runs")
tool_calls = pd.read_excel(RESULTS_PATH, sheet_name="Tool_Calls")
runs["order_id"] = runs["order_id"].astype(str)
tool_calls["order_id"] = tool_calls["order_id"].astype(str)

flagged_order_ids = {order_id for order_id, *_ in FLAGGED}
prompts = {}
with open(JSONL_PATH) as f:
    for line in f:
        rec = json.loads(line)
        if rec.get("model") == "gemma4:e4b" and rec.get("tool_mode") == "tools" \
                and str(rec.get("order_id")) in flagged_order_ids:
            prompts[str(rec["order_id"])] = rec["prompt"]
missing = flagged_order_ids - prompts.keys()
if missing:
    raise SystemExit(f"no prompt found in {JSONL_PATH} for order_id(s): {missing}")

rows = []
transcripts = []
for order_id, run_index, match_type, why in FLAGGED:
    run = runs[(runs["model"] == "gemma4:e4b") & (runs["tool_mode"] == "tools")
               & (runs["order_id"] == order_id) & (runs["run_index"] == run_index)]
    if len(run) != 1:
        raise SystemExit(f"expected exactly 1 run for order_id={order_id!r} run_index={run_index}, got {len(run)}")
    run = run.iloc[0]

    calls = tool_calls[(tool_calls["model"] == "gemma4:e4b") & (tool_calls["tool_mode"] == "tools")
                        & (tool_calls["order_id"] == order_id) & (tool_calls["run_index"] == run_index)]
    calls = calls.sort_values("round")
    calls_text = "\n".join(
        f"round {c.round} | {c.tool_name}({c.arguments}) -> {c.result}"
        for c in calls.itertuples()
    )

    rows.append({
        "match_type": match_type,
        "case_index": run["case_index"],
        "label": run["label"],
        "queue_length_bucket": run["queue_length_bucket"],
        "order_id": order_id,
        "run_index": run_index,
        "correct_position": run["correct_position"],
        "model_position": run["model_position"],
        "verdict": run["verdict"],
        "n_rounds": run["n_rounds"],
        "why_flagged": why,
        "prompt": prompts[order_id],
        "tool_calls": calls_text,
        "thinking": run["thinking"],
        "content": run["content"],
    })

cases = pd.DataFrame(rows)
cases = cases.sort_values(["match_type", "case_index"], ascending=[False, True])

print(f"{len(cases)} flagged cases ({(cases['match_type'] == 'clean').sum()} clean, "
      f"{(cases['match_type'] == 'muddled').sum()} muddled)")
print(cases[["match_type", "case_index", "label", "correct_position", "model_position"]].to_string(index=False))

notes = pd.DataFrame({"note": [
    "gemma4:e4b, tool-calling mode, run_index == 1 (first/original run per case) only.",
    "'clean': the final answer is literally the position of whichever item just beat "
    "new_order in a pairwise compare_priority/compare_wait_time call, instead of "
    "new_order's own rank.",
    "'muddled': same root cause -- the thinking text explicitly states new_order is "
    "NOT Position 1 -- but instead of substituting the winner's position it guesses "
    "an arbitrary small number under the 'reply with ONLY Position: N' formatting "
    "instruction.",
    "11 of 16 non-PASS gemma4:e4b/tools run_index==1 cases match this pattern (4.4% "
    "of all 250 tools-mode first-run cases). It does not occur in no-tools mode -- "
    "that mode's 2 run_index==1 failures are a separate empty-queue-hallucination bug.",
    "See nni_llm_queue_slot_consistency_gemma_wrong_entity_cases.py for the manual "
    "case list (order_id, run_index) and the exact quote that justified each flag.",
    "The 'prompt' column is pulled from the raw JSONL checkpoint (nni_llm_queue_slot_"
    "consistency_raw.jsonl), not from this workbook's own Runs sheet -- the literal "
    "prompt text isn't carried into the assembled results xlsx, only thinking/content.",
]})

with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
    notes.to_excel(writer, sheet_name="Notes", index=False)
    cases.to_excel(writer, sheet_name="Cases", index=False)
    for sheet in writer.sheets.values():
        auto_col_width(sheet)
    # thinking/content/tool_calls are long free text -- wrap and widen those columns specifically
    ws = writer.sheets["Cases"]
    from openpyxl.styles import Alignment
    wide_cols = {"prompt": 70, "tool_calls": 60, "thinking": 90, "content": 30, "why_flagged": 50}
    header = [c.value for c in ws[1]]
    for col_name, width in wide_cols.items():
        col_idx = header.index(col_name) + 1
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = width
        for cell in ws[letter][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

print(f"\nWrote {OUT_PATH}")
